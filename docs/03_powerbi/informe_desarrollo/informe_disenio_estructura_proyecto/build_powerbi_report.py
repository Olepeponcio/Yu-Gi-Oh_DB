from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
OUT_PATH = BASE_DIR / "informe_analisis_powerbi_yugioh.docx"
PBIX_PATH = BASE_DIR.parents[2] / "power_bi" / "informes" / "analisis_yugioh_db.pbix"
VIDEO_DIAGNOSTICO = (
    BASE_DIR.parent
    / "video_proceso_analaisis"
    / "analisis_diagnostico"
    / "video_medidas_03_analisis_diagnostico.mp4"
)


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_GRAY = "F2F4F7"
BORDER = "D9E2EC"
MUTED = "666666"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def style_table(table, widths: list[int] | None = None, header: bool = True) -> None:
    set_table_width(table)
    table.autofit = False
    if table.rows:
        mark_header_row(table.rows[0])
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[cell_idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if header and row_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Actualizar campo en Word"
    fld_sep.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    title.paragraph_format.space_after = Pt(6)


def add_footer(section) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Informe Power BI - Yu-Gi-Oh! | Pagina ")
    add_field(footer, "PAGE")


def add_metadata_table(doc: Document) -> None:
    rows = [
        ("Proyecto", "Proyecto SQL DB Yu-Gi-Oh"),
        ("Documento", "Informe de diseño y documentacion del analisis en Power BI"),
        ("Autor", "Pepin"),
        ("Fecha", "Julio 2026"),
        ("Version", "0.2 - informe alineado con PBIX y videos de proceso"),
        (
            "Estado",
            "Documento vivo: ya registra paginas, visuales y medidas detectadas en el PBIX.",
        ),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    for idx, (label, value) in enumerate(rows):
        table.cell(idx, 0).text = label
        table.cell(idx, 1).text = value
    style_table(table, widths=[2500, 6860], header=False)
    for row in table.rows:
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)


def add_status(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = text
    style_table(table, widths=[9360], header=False)
    set_cell_shading(table.cell(0, 0), "F8FAFC")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_kv_table(
    doc: Document, rows: list[tuple[str, str]], widths=(2600, 6760)
) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Elemento"
    table.cell(0, 1).text = "Detalle"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    style_table(table, widths=list(widths), header=True)


def add_matrix(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, widths=widths, header=True)


def build_doc() -> None:
    doc = Document()
    configure_styles(doc)
    add_footer(doc.sections[0])
    set_update_fields(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Analisis del Mercado de Cartas Yu-Gi-Oh!")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Informe de diseño y documentación del análisis en Power BI")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph()
    add_metadata_table(doc)
    doc.add_paragraph()
    add_status(
        doc,
        "Proposito del documento: acompanar al dashboard de Power BI. El panel responde que ocurre; "
        "este informe documenta por que se analiza, como se ha construido el modelo y que conclusiones "
        "se podran defender cuando existan visuales y medidas validadas.",
    )

    doc.add_page_break()
    doc.add_heading("Indice", level=1)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph(
        "Nota: en Word, actualizar campos para regenerar el indice cuando el documento evolucione."
    )

    doc.add_page_break()
    doc.add_heading("1. Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Este informe documenta el diseño del analisis en Power BI para un proyecto de datos de cartas "
        "Yu-Gi-Oh. El origen operativo es una extraccion desde la API publica YGOPRODeck, normalizada "
        "mediante un ETL en Python y cargada en MySQL bajo el schema yugioh_db."
    )
    doc.add_paragraph(
        "La fase actual deja preparado el marco de consumo analitico: dimensiones y hechos base en formato `vw_`. "
        "Power BI consume esas vistas sin modificar las tablas madre."
    )
    add_status(
        doc,
        "Estado actual: el PBIX analisis_yugioh_db.pbix contiene indice, vista general, analisis descriptivo, "
        "analisis diagnostico, paginas 04-06 aun sin desarrollo funcional y tooltip de informacion de carta. "
        "El informe documenta estos avances y deja pendientes resultados cuantitativos cerrados.",
    )

    doc.add_heading("2. Objetivos del proyecto", level=1)
    add_bullets(
        doc,
        [
            "Analizar el catalogo de cartas Yu-Gi-Oh disponible desde YGOPRODeck.",
            "Construir un modelo Power BI trazable desde MySQL y no desde reglas improvisadas en visuales.",
            "Describir precios actuales por carta, marketplace y moneda.",
            "Analizar sets, rarezas, apariciones y precios desde hechos base.",
            "Preparar una base para recomendaciones prescriptivas revisables, no basadas en rankings aislados.",
        ],
    )

    doc.add_heading("3. Descripcion de los datos", level=1)
    add_kv_table(
        doc,
        [
            ("Origen externo", "API publica YGOPRODeck."),
            ("Origen interno para Power BI", "MySQL schema yugioh_db."),
            (
                "Extraccion",
                "src/api/ygoprodeck_client.py guarda data/raw/cardinfo_latest.json.",
            ),
            ("Carga", "ETL Python hacia tablas madre definidas en sql/schema.sql."),
            (
                "Numero de registros",
                "Pendiente: depende de la ejecucion local del ETL y de la fecha de extraccion.",
            ),
            ("Numero de tablas madre", "10 tablas principales."),
            (
                "Historico",
                "card_price_history inserta snapshots de precios en cada carga real del ETL.",
            ),
        ],
    )

    doc.add_heading("3.1 Tablas madre", level=2)
    add_matrix(
        doc,
        ["Tabla", "Funcion", "Grano principal"],
        [
            ["cards", "Catalogo base de cartas.", "1 carta"],
            ["sets", "Catalogo de sets.", "1 set"],
            [
                "rarities",
                "Catalogo tecnico de rarezas por codigo.",
                "1 set_code + rareza + codigo",
            ],
            [
                "card_sets",
                "Apariciones de cartas en sets.",
                "1 carta + 1 set/codigo + 1 rareza",
            ],
            ["card_images", "Imagenes asociadas.", "1 imagen"],
            ["card_prices", "Precios actuales por marketplace.", "1 carta"],
            [
                "card_price_history",
                "Snapshots historicos de precios.",
                "1 carta + 1 snapshot",
            ],
            ["card_banlist", "Estado de banlist.", "1 carta"],
            ["card_typelines", "Typelines por carta.", "1 carta + 1 typeline"],
            [
                "card_linkmarkers",
                "Marcadores Link por carta.",
                "1 carta + 1 linkmarker",
            ],
        ],
        [1900, 4300, 3160],
    )

    doc.add_heading("4. Preparacion de datos", level=1)
    doc.add_paragraph(
        "La preparacion se realiza antes de Power BI. El programa Python extrae el JSON, conserva una copia raw, "
        "normaliza dominios y carga MySQL. Power BI debe consumir consultas documentadas, no reconstruir la logica "
        "de negocio desde tablas crudas sin control."
    )
    add_numbered(
        doc,
        [
            "Descargar JSON o leer un JSON local.",
            "Normalizar cartas, sets, rarezas, apariciones, precios, imagenes, banlist y relaciones.",
            "Validar claves y campos requeridos.",
            "Insertar o actualizar tablas madre.",
            "Registrar snapshot de precios en card_price_history.",
            "Publicar vistas SQL de consumo para Power BI.",
        ],
    )

    doc.add_heading("4.1 Reglas criticas de preparacion", level=2)
    add_bullets(
        doc,
        [
            "No mezclar monedas sin segmentacion o conversion explicita.",
            "cardmarket_price llega en EUR; tcgplayer_price, ebay_price, amazon_price y coolstuffinc_price llegan en USD.",
            "set_price pertenece a card_sets; no es un precio propio de la rareza.",
            "Los hechos base cargados en Power BI son precios actuales, apariciones carta-set-rareza y variacion historica.",
            "Los rankings, outliers y resumenes se calculan como medidas o filtros desde hechos base.",
        ],
    )

    doc.add_heading("4.2 Procesos aplicados para el informe Power BI", level=2)
    add_matrix(
        doc,
        ["Proceso", "Aplicacion", "Criterio vigente"],
        [
            [
                "Carga y transformacion ETL",
                "Datos extraidos, transformados y cargados en MySQL desde el flujo Python.",
                "Power BI consume las vistas SQL publicadas; no modifica tablas madre.",
            ],
            [
                "Nulos en atributos numericos de carta",
                "`atk`, `def` y `link_value` se interpretan como 0 cuando no aplican; `scale` se interpreta como -1.",
                "Evita mezclar nulos tecnicos con valores analizables en segmentaciones y tarjetas.",
            ],
            [
                "Variacion porcentual de precio",
                "`vw_fact_card_price_variation_predictive.price_change_pct` conserva `NULL` cuando no hay base comparable.",
                "Las metricas y visuales aplican contexto de filtro para excluir valores vacios cuando el calculo lo requiera.",
            ],
            [
                "Filtrado aplicado en Power BI",
                "Las filas con `price_change_pct` vacio se excluyen en pasos aplicados o filtros del visual cuando se analizan variaciones.",
                "La vista SQL conserva el dato original; el filtro pertenece al contexto analitico.",
            ],
        ],
        [2500, 4300, 2560],
    )

    doc.add_heading("5. Arquitectura y modelo de datos", level=1)
    doc.add_paragraph(
        "El modelo recomendado para Power BI es una constelacion simplificada de hechos. Las dimensiones filtran "
        "hechos con direccion unica 1 -> *. No se recomienda relacionar hechos entre si, porque aumenta el riesgo "
        "de filtros ambiguos y dobles conteos."
    )
    add_status(
        doc,
        "Cambio de mantenimiento: este constructor ya no genera ni inserta modelo_relacional_powerbi.png. "
        "El modelo queda documentado mediante tablas para evitar artefactos graficos duplicados.",
    )

    add_matrix(
        doc,
        ["Grupo", "Tablas / vistas", "Funcion en Power BI"],
        [
            [
                "Dimensiones descriptivas",
                "vw_dim_cards_descriptive, vw_dim_sets_descriptive, vw_dim_rarities_descriptive",
                "Aportan etiquetas, atributos y agrupaciones para filtrar hechos.",
            ],
            [
                "Dimensiones de precio",
                "vw_dim_marketplaces_descriptive, vw_dim_currencies_descriptive",
                "Controlan fuente y moneda antes de comparar precios.",
            ],
            [
                "Hechos actuales",
                "vw_fact_card_prices_descriptive",
                "Calcula precio por carta, marketplace y moneda.",
            ],
            [
                "Hecho puente",
                "vw_fact_card_set_appearances",
                "Explica apariciones carta-set-rareza y presencia historica.",
            ],
            [
                "Hecho historico",
                "vw_fact_card_price_variation_predictive",
                "Soporta variacion temporal cuando hay snapshots suficientes.",
            ],
        ],
        [2300, 3600, 3460],
    )

    doc.add_heading("5.1 Relaciones recomendadas", level=2)
    add_bullets(
        doc,
        [
            "vw_dim_cards_descriptive 1 -> * vw_fact_card_prices_descriptive.",
            "vw_dim_cards_descriptive 1 -> * vw_fact_card_set_appearances.",
            "vw_dim_cards_descriptive 1 -> * vw_fact_card_price_variation_predictive.",
            "vw_dim_sets_descriptive 1 -> * vw_fact_card_set_appearances.",
            "vw_dim_rarities_descriptive 1 -> * vw_fact_card_set_appearances.",
            "vw_dim_marketplaces_descriptive 1 -> * hechos de precios y variacion.",
            "vw_dim_currencies_descriptive 1 -> * hechos de precios y variacion.",
            "vw_dim_snapshots_descriptive 1 -> * vw_fact_card_price_variation_predictive.",
        ],
    )

    doc.add_heading("6. Vistas de consumo para Power BI", level=1)
    add_matrix(
        doc,
        ["Vista", "Tipo", "Uso recomendado"],
        [
            ["vw_dim_cards_descriptive", "Dimension", "Catalogo base de cartas."],
            ["vw_dim_sets_descriptive", "Dimension", "Catalogo de sets."],
            [
                "vw_dim_rarities_descriptive",
                "Dimension",
                "Catalogo tecnico de rarezas.",
            ],
            [
                "vw_dim_marketplaces_descriptive",
                "Dimension",
                "Segmentacion de fuentes de precio.",
            ],
            ["vw_dim_currencies_descriptive", "Dimension", "Segmentacion por moneda."],
            [
                "vw_dim_snapshots_descriptive",
                "Dimension",
                "Calendario real de historico disponible.",
            ],
            [
                "vw_fact_card_prices_descriptive",
                "Hecho",
                "Precios actuales en formato largo.",
            ],
            [
                "vw_fact_card_set_appearances",
                "Hecho puente",
                "Apariciones carta-set-rareza y precio de set.",
            ],
            [
                "vw_fact_card_price_variation_predictive",
                "Hecho historico",
                "Variacion temporal por carta, marketplace y moneda.",
            ],
        ],
        [3950, 1950, 3460],
    )

    doc.add_heading("7. Metodologia analitica", level=1)
    add_numbered(
        doc,
        [
            "Analisis descriptivo: entender que existe, como se distribuye y que cobertura tienen los datos.",
            "Analisis diagnostico: explicar por que ciertos precios, rarezas o cartas destacan.",
            "Analisis predictivo: estudiar variacion temporal solo si hay snapshots suficientes.",
            "Analisis prescriptivo: convertir criterios validados en decisiones revisables.",
        ],
    )

    doc.add_heading("8. Analisis descriptivo", level=1)
    doc.add_paragraph(
        "La pagina 02_Analisis_descriptivo ya contiene visuales orientados a precios por marketplace, "
        "ranking de cartas y valor por set. El patron de lectura debe separar precio actual, fuente y moneda."
    )
    add_matrix(
        doc,
        ["Visual detectado en PBIX", "Campos / medidas", "Lectura"],
        [
            [
                "Tabla: Precio medio de mercado USD",
                "name, Precio Amazon USD, Precio CoolStuffInc USD, Precio eBay USD, Precio TCGplayer USD, Precio medio USD",
                "Comparacion tabular de precios por carta y fuente USD.",
            ],
            [
                "Barras: Precio medio global por Marketplace USD",
                "marketplace_name, Sum(price)",
                "Resumen por fuente; requiere mantener filtro de moneda.",
            ],
            [
                "Barras agrupadas: cartas mayor precio medio marketplace",
                "card_name, marketplace, Precio medio marketplace USD",
                "Ranking comparativo por carta y marketplace.",
            ],
            [
                "Barras: valor de mercado por set",
                "set_name, Valor Mercado Set",
                "Lectura de concentracion de valor por set.",
            ],
        ],
        [3000, 4060, 2300],
    )

    doc.add_heading("9. Analisis diagnostico", level=1)
    add_matrix(
        doc,
        ["Visual detectado en PBIX", "Campos / medidas", "Criterio de interpretacion"],
        [
            [
                "Barras: Precio Mediana por rareza",
                "rarity_name, Precio Mediano por Rareza",
                "Diagnostica rarezas asociadas a precios mas altos; usar mediana reduce impacto de outliers.",
            ],
            [
                "Slicer: marketplace",
                "marketplace",
                "Evita mezclar fuentes; en la captura de trabajo se uso Amazon como contexto.",
            ],
            [
                "Barras: Sets Distintos por Carta",
                "card_name, Sets Distintos por Carta; tooltip Apariciones en Sets",
                "Explica presencia, reimpresion o disponibilidad. No equivale a recomendacion de compra.",
            ],
        ],
        [3000, 3860, 2500],
    )
    add_status(
        doc,
        "Material complementario generado: video_proceso_analaisis/analisis_diagnostico/"
        "video_medidas_03_analisis_diagnostico.mp4 explica el razonamiento tabla -> columna -> relacion -> medida -> visual.",
    )

    doc.add_heading("10. Analisis predictivo", level=1)
    doc.add_paragraph(
        "El analisis predictivo queda condicionado por el numero de snapshots reales en card_price_history. "
        "Sin historico suficiente no debe hablarse de tendencia; solo de disponibilidad de datos."
    )
    add_bullets(
        doc,
        [
            "vw_dim_snapshots_descriptive: permite validar fechas de snapshot disponibles.",
            "vw_fact_card_price_variation_predictive: compara precio actual contra snapshot anterior por carta, marketplace y moneda.",
            "Requiere conservar moneda y marketplace durante todo el analisis.",
        ],
    )

    doc.add_heading("11. Analisis prescriptivo", level=1)
    doc.add_paragraph(
        "La parte prescriptiva debe esperar a que descriptivo, diagnostico y predictivo esten validados. "
        "No se deben convertir rankings aislados en recomendaciones."
    )
    add_matrix(
        doc,
        ["Decision futura", "Criterio de avance", "Estado"],
        [
            [
                "Seleccionar carta principal potencial",
                "Combinar valor, presencia y legalidad.",
                "Pendiente",
            ],
            [
                "Seleccionar carta complementaria",
                "Precio moderado y coherencia tematica.",
                "Pendiente",
            ],
            [
                "Marcar carta para revision",
                "Outlier, moneda mezclada o dato incompleto.",
                "Pendiente",
            ],
        ],
        [3000, 4300, 2060],
    )

    doc.add_heading("12. Paginas previstas del dashboard", level=1)
    add_matrix(
        doc,
        ["Ordinal PBIX", "Pagina", "Estado detectado"],
        [
            [
                "0",
                "Indice",
                "Pagina de navegacion con botones y formas.",
            ],
            [
                "1",
                "01_vista_general",
                "Contiene tarjetas de volumen de cartas, sets, snapshots, ultimo snapshot y tablas de marketplaces/rarezas.",
            ],
            [
                "2",
                "02_Analisis_descriptivo",
                "Contiene tabla de precio medio USD, barras por marketplace, ranking por carta-marketplace y valor por set.",
            ],
            [
                "3",
                "03_Análisis_diagnostico",
                "Contiene precio mediano por rareza, slicer marketplace y sets distintos por carta con tooltip de apariciones.",
            ],
            [
                "4",
                "04_",
                "Solo boton de navegacion; pendiente de desarrollo analitico.",
            ],
            [
                "5",
                "05_",
                "Solo boton de navegacion; pendiente de desarrollo analitico.",
            ],
            [
                "6",
                "06_",
                "Solo boton de navegacion; pendiente de desarrollo analitico.",
            ],
            [
                "7",
                "tooltip_01_cards_information",
                "Tooltip con imagen, nombre de carta y set asociado.",
            ],
        ],
        [1300, 2800, 5260],
    )

    doc.add_heading("13. Conclusiones iniciales", level=1)
    add_bullets(
        doc,
        [
            "El proyecto ya separa responsabilidades: Python carga, MySQL conserva y Power BI consume vistas documentadas.",
            "El modelo BI debe tratarse como constelacion de hechos con dimensiones compartidas.",
            "El precio debe analizarse siempre con moneda declarada; mezclar EUR y USD invalida comparaciones directas.",
            "Los outliers son una lista de revision, no una conclusion automatica.",
            "La fase prescriptiva requiere reglas validadas antes de emitir recomendaciones.",
        ],
    )

    doc.add_heading("14. Limitaciones", level=1)
    add_bullets(
        doc,
        [
            "No se han incluido todavia resultados cuantitativos de una ejecucion local concreta.",
            "El PBIX permite detectar visuales y nombres de medidas, pero las formulas DAX deben validarse en Power BI Desktop antes de tratarlas como definitivas.",
            "Los precios proceden de campos de marketplaces y pueden cambiar con cada actualizacion.",
            "No se consideran costes de envio, estado fisico de la carta, ventas privadas ni liquidez real.",
            "El historico depende de ejecuciones reales del ETL; sin suficientes snapshots no hay tendencia robusta.",
        ],
    )

    doc.add_heading("15. Trabajo futuro", level=1)
    add_bullets(
        doc,
        [
            "Renombrar y desarrollar las paginas 04, 05 y 06 segun su finalidad analitica.",
            "Documentar formulas DAX definitivas desde Power BI Desktop.",
            "Insertar capturas de cada visual validado cuando el dashboard estabilice diseno y datos.",
            "Validar nulos, duplicados y cobertura antes de interpretar hallazgos.",
            "Ampliar el analisis prescriptivo con criterios revisables y no automaticos.",
        ],
    )

    doc.add_heading("16. Anexos", level=1)
    doc.add_heading("16.1 Diccionario operativo minimo", level=2)
    add_kv_table(
        doc,
        [
            ("Carta", "Entidad base identificada por cards.card_id."),
            ("Aparicion", "Fila de card_sets que combina carta, set/codigo y rareza."),
            (
                "Marketplace",
                "Fuente de precio: Cardmarket, TCGPlayer, eBay, Amazon o CoolStuffInc.",
            ),
            (
                "Snapshot",
                "Foto historica de precios insertada por una ejecucion real del ETL.",
            ),
            (
                "Outlier",
                "Precio candidato a revision; no implica oportunidad ni error por si mismo.",
            ),
        ],
    )

    doc.add_heading("16.2 Registro de medidas DAX", level=2)
    add_matrix(
        doc,
        ["Medida", "Uso detectado en PBIX", "Estado"],
        [
            ["Volumen de cartas", "Tarjeta en 01_vista_general.", "Detectada"],
            ["Volumen de sets", "Tarjeta en 01_vista_general.", "Detectada"],
            ["Volumen de rarezas", "Tabla de rarezas en 01_vista_general.", "Detectada"],
            ["Numero rareras nombres", "Tabla de rarezas en 01_vista_general.", "Detectada; revisar nombre."],
            ["Total snapshots", "Tarjeta en 01_vista_general.", "Detectada"],
            ["Ultimo snapshot", "Tarjeta en 01_vista_general.", "Detectada"],
            ["Precio Amazon USD", "Tabla de precio medio de mercado USD.", "Detectada"],
            ["Precio CoolStuffInc USD", "Tabla de precio medio de mercado USD.", "Detectada"],
            ["Precio eBay USD", "Tabla de precio medio de mercado USD.", "Detectada"],
            ["Precio TCGplayer USD", "Tabla de precio medio de mercado USD.", "Detectada"],
            ["Precio medio USD", "Tabla descriptiva de precios.", "Detectada"],
            ["Precio medio marketplace USD", "Ranking carta-marketplace.", "Detectada"],
            ["Valor Mercado Set", "Barra por set en analisis descriptivo.", "Detectada"],
            ["Precio Mediano por Rareza", "Barra diagnostica por rarity_name.", "Detectada; validar puente rareza-precio."],
            ["Sets Distintos por Carta", "Barra diagnostica por card_name.", "Detectada"],
            ["Apariciones en Sets", "Tooltip del ranking de cartas.", "Detectada"],
        ],
        [2600, 4760, 2000],
    )

    doc.add_heading("16.3 Registro de mantenimiento del informe", level=2)
    add_matrix(
        doc,
        ["Elemento", "Responsable", "Estado"],
        [["Modelo relacional Power BI", "Proyecto", "Vigente"]],
        [1800, 1800, 5760],
    )

    doc.save(OUT_PATH)


if __name__ == "__main__":
    try:
        build_doc()
        print(OUT_PATH)
    except PermissionError:
        OUT_PATH = BASE_DIR / "informe_analisis_powerbi_yugioh_actualizado.docx"
        build_doc()
        print(OUT_PATH)
